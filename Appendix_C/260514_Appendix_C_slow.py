#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
One-file comparison:
1) ADA-L-based continuous-time trajectory optimization (Adam + optional L-BFGS-B)
2) Joint-space quintic polynomial baseline

Both are evaluated with the same benchmark metrics:
- final_position_error
- joint_path_length
- ee_path_length
- integrated_squared_velocity
- integrated_squared_acceleration
- integrated_squared_jerk
- mean_squared_velocity
- mean_squared_acceleration
- mean_squared_jerk
- max_velocity_norm
- max_acceleration_norm
- max_jerk_norm
- terminal_velocity_norm
- terminal_acceleration_norm
- solve_time

Notes
-----
- Quintic baseline here is the classical minimum-jerk-like joint interpolation
  with boundary conditions:
      q(0)=q0, q(T)=qT, qdot(0)=qdot(T)=0, qddot(0)=qddot(T)=0
- Since the user gives a Cartesian target, the quintic baseline first finds qT
  by solving a terminal IK-style optimization with TensorFlow Adam.
- Then it generates the closed-form quintic trajectory from q0 to qT.
- ADA-L planner is also included so both methods can be compared directly in one file.

Requirements
------------
pip install tensorflow matplotlib sympy scipy numpy
"""

import os
import time
import numpy as np
import tensorflow as tf
import matplotlib
# matplotlib.rcParams['animation.ffmpeg_path'] = r'C:\Users\user\anaconda3\Library\bin\ffmpeg.exe'
import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator, FormatStrFormatter, FuncFormatter
import sympy as sp
import scipy.optimize

import json
import csv
from pathlib import Path
from datetime import datetime

os.environ["TF_CPP_MIN_LOG_LEVEL"] = "2"

DTYPE = tf.float32
np.random.seed(0)
tf.random.set_seed(0)

fontsize = 14

plt.rcParams["font.family"] = "Arial"
plt.rcParams["font.size"] = fontsize
plt.rcParams["axes.titlesize"] = fontsize
plt.rcParams["axes.labelsize"] = fontsize
plt.rcParams["xtick.labelsize"] = fontsize
plt.rcParams["ytick.labelsize"] = fontsize
plt.rcParams["legend.fontsize"] = fontsize

# ============================================================
# Extra utilities for saving figures and generating report
# ============================================================
def ensure_dir(path):
    Path(path).mkdir(parents=True, exist_ok=True)


def save_json(data, path):
    def convert(v):
        if isinstance(v, np.ndarray):
            return v.tolist()
        if isinstance(v, (np.float32, np.float64, np.float16)):
            return float(v)
        if isinstance(v, (np.int32, np.int64, np.int16)):
            return int(v)
        return v

    data2 = {k: convert(v) for k, v in data.items()}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data2, f, indent=2, ensure_ascii=False)


def save_comparison_csv(metrics_a, metrics_b, path, name_a="ADA-L", name_b="Quintic"):
    keys = [
        "final_position_error",
        "joint_path_length",
        "ee_path_length",
        "integrated_squared_velocity",
        "integrated_squared_acceleration",
        "integrated_squared_jerk",
        "mean_squared_velocity",
        "mean_squared_acceleration",
        "mean_squared_jerk",
        "max_velocity_norm",
        "max_acceleration_norm",
        "max_jerk_norm",
        "terminal_velocity_norm",
        "terminal_acceleration_norm",
        "solve_time",
    ]

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["metric", name_a, name_b, f"{name_a}/{name_b}", f"{name_b}/{name_a}"])
        for k in keys:
            a = float(metrics_a[k])
            b = float(metrics_b[k])
            ratio_ab = a / b if abs(b) > 1e-15 else np.nan
            ratio_ba = b / a if abs(a) > 1e-15 else np.nan
            writer.writerow([k, a, b, ratio_ab, ratio_ba])

def save_joint_trajectory_csv(t, q, path):
    """Save joint trajectory CSV with theta_i = q + theta_offset (robot convention)."""
    t = np.asarray(t, dtype=np.float32).reshape(-1)
    q = np.asarray(q, dtype=np.float32)

    if q.ndim != 2 or q.shape[1] != 6:
        raise ValueError("q must have shape (N, 6)")
    if q.shape[0] != t.shape[0]:
        raise ValueError("len(t) must match q.shape[0]")

    theta_offset = np.array(UR5eKinematics().theta_offset, dtype=np.float32)
    q_robot = q + theta_offset[None, :]
    header = ["time", "q1", "q2", "q3", "q4", "q5", "q6"]

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for i in range(len(t)):
            writer.writerow([float(t[i])] + [float(v) for v in q_robot[i]])

def metric_better_text(metric_name, a, b, smaller_is_better=True):
    if smaller_is_better:
        if a < b:
            return f"ADA-L better ({a:.4e} < {b:.4e})"
        elif a > b:
            return f"Quintic better ({b:.4e} < {a:.4e})"
        else:
            return "Tie"
    else:
        if a > b:
            return f"ADA-L better ({a:.4e} > {b:.4e})"
        elif a < b:
            return f"Quintic better ({b:.4e} > {a:.4e})"
        else:
            return "Tie"


def generate_metric_summary_text(metrics_adal, metrics_quintic):
    lines = []
    lines.append("## Automatic comparison summary\n")

    smaller_better_metrics = [
        "final_position_error",
        "joint_path_length",
        "ee_path_length",
        "integrated_squared_velocity",
        "integrated_squared_acceleration",
        "integrated_squared_jerk",
        "mean_squared_velocity",
        "mean_squared_acceleration",
        "mean_squared_jerk",
        "max_velocity_norm",
        "max_acceleration_norm",
        "max_jerk_norm",
        "terminal_velocity_norm",
        "terminal_acceleration_norm",
        "solve_time",
    ]

    for k in smaller_better_metrics:
        a = float(metrics_adal[k])
        b = float(metrics_quintic[k])
        lines.append(f"- **{k}**: {metric_better_text(k, a, b, smaller_is_better=True)}")

    return "\n".join(lines)


def write_markdown_report(
    report_path,
    run_config,
    metrics_adal,
    metrics_quintic,
    fig_paths,
    adal_final_ee,
    quintic_final_ee,
    target_xyz,
):
    summary_text = generate_metric_summary_text(metrics_adal, metrics_quintic)

    lines = []
    lines.append("# ADA-L vs Quintic Trajectory Comparison Report\n")
    lines.append(f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")

    lines.append("## Problem setup\n")
    lines.append(f"- q0: `{np.array2string(run_config['q0'], precision=5)}`")
    lines.append(f"- target_xyz: `{np.array2string(run_config['target_xyz'], precision=5)}`")
    lines.append(f"- T_final: `{run_config['T_final']}`")
    lines.append(f"- Nt: `{run_config['Nt']}`\n")

    lines.append("## Final end-effector positions\n")
    lines.append(f"- ADA-L final EE: `{np.array2string(adal_final_ee, precision=6)}`")
    lines.append(f"- Quintic final EE: `{np.array2string(quintic_final_ee, precision=6)}`")
    lines.append(f"- Target: `{np.array2string(target_xyz, precision=6)}`\n")

    lines.append("## Scalar metrics table\n")
    lines.append("| Metric | ADA-L | Quintic |")
    lines.append("|---|---:|---:|")

    keys = [
        "final_position_error",
        "joint_path_length",
        "ee_path_length",
        "integrated_squared_velocity",
        "integrated_squared_acceleration",
        "integrated_squared_jerk",
        "mean_squared_velocity",
        "mean_squared_acceleration",
        "mean_squared_jerk",
        "max_velocity_norm",
        "max_acceleration_norm",
        "max_jerk_norm",
        "terminal_velocity_norm",
        "terminal_acceleration_norm",
        "solve_time",
    ]
    for k in keys:
        lines.append(f"| {k} | {float(metrics_adal[k]):.6e} | {float(metrics_quintic[k]):.6e} |")
    lines.append("")

    lines.append(summary_text)
    lines.append("")

    lines.append("## Per-joint extrema\n")
    lines.append(f"- ADA-L max abs velocity each joint: `{np.array2string(metrics_adal['max_abs_velocity_each_joint'], precision=5)}`")
    lines.append(f"- Quintic max abs velocity each joint: `{np.array2string(metrics_quintic['max_abs_velocity_each_joint'], precision=5)}`")
    lines.append(f"- ADA-L max abs acceleration each joint: `{np.array2string(metrics_adal['max_abs_acceleration_each_joint'], precision=5)}`")
    lines.append(f"- Quintic max abs acceleration each joint: `{np.array2string(metrics_quintic['max_abs_acceleration_each_joint'], precision=5)}`")
    lines.append(f"- ADA-L max abs jerk each joint: `{np.array2string(metrics_adal['max_abs_jerk_each_joint'], precision=5)}`")
    lines.append(f"- Quintic max abs jerk each joint: `{np.array2string(metrics_quintic['max_abs_jerk_each_joint'], precision=5)}`\n")

    lines.append("## Figures\n")
    for title, rel_path in fig_paths:
        lines.append(f"### {title}\n")
        lines.append(f"![{title}]({rel_path})\n")

    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def try_write_docx_report(
    docx_path,
    run_config,
    metrics_adal,
    metrics_quintic,
    fig_abs_paths,
    adal_final_ee,
    quintic_final_ee,
    target_xyz,
):
    try:
        from docx import Document
        from docx.shared import Inches
    except Exception as e:
        print(f"[report] python-docx unavailable, skipping DOCX generation: {e}")
        return False

    doc = Document()
    doc.add_heading("ADA-L vs Quintic Trajectory Comparison Report", level=0)

    doc.add_heading("Problem setup", level=1)
    doc.add_paragraph(f"q0: {np.array2string(run_config['q0'], precision=5)}")
    doc.add_paragraph(f"target_xyz: {np.array2string(run_config['target_xyz'], precision=5)}")
    doc.add_paragraph(f"T_final: {run_config['T_final']}")
    doc.add_paragraph(f"Nt: {run_config['Nt']}")

    doc.add_heading("Final end-effector positions", level=1)
    doc.add_paragraph(f"ADA-L final EE: {np.array2string(adal_final_ee, precision=6)}")
    doc.add_paragraph(f"Quintic final EE: {np.array2string(quintic_final_ee, precision=6)}")
    doc.add_paragraph(f"Target: {np.array2string(target_xyz, precision=6)}")

    doc.add_heading("Scalar metrics", level=1)
    keys = [
        "final_position_error",
        "joint_path_length",
        "ee_path_length",
        "integrated_squared_velocity",
        "integrated_squared_acceleration",
        "integrated_squared_jerk",
        "mean_squared_velocity",
        "mean_squared_acceleration",
        "mean_squared_jerk",
        "max_velocity_norm",
        "max_acceleration_norm",
        "max_jerk_norm",
        "terminal_velocity_norm",
        "terminal_acceleration_norm",
        "solve_time",
    ]

    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "ADA-L"
    hdr[2].text = "Quintic"

    for k in keys:
        row = table.add_row().cells
        row[0].text = k
        row[1].text = f"{float(metrics_adal[k]):.6e}"
        row[2].text = f"{float(metrics_quintic[k]):.6e}"

    doc.add_heading("Figures", level=1)
    for title, abs_path in fig_abs_paths:
        doc.add_heading(title, level=2)
        doc.add_picture(str(abs_path), width=Inches(6.3))

    doc.save(docx_path)
    print(f"[report] DOCX saved to: {docx_path}")
    return True

# ============================================================
# 1) Symbolic helpers for Legendre anti-derivative basis
# ============================================================
def sympy_poly_to_numpy_coeffs(expr, x_sym, dtype=np.float32):
    poly = sp.Poly(sp.expand(expr), x_sym)
    coeff_dict = poly.as_dict()
    max_deg = poly.degree()

    coeffs = np.zeros((max_deg + 1,), dtype=np.float64)
    for k, v in coeff_dict.items():
        coeffs[k[0]] = float(v)
    return coeffs.astype(dtype)


def build_legendre_symbolic_tables(max_order=6, dtype=np.float32):
    x = sp.symbols("x", real=True)

    P_coeffs = {}
    I1_coeffs = {}
    I2_coeffs = {}
    I3_coeffs = {}

    for n in range(max_order + 1):
        Pn = sp.legendre(n, x)

        I1 = sp.integrate(Pn, x)
        I1 = sp.expand(I1 - I1.subs(x, -1))

        I2 = sp.integrate(I1, x)
        I2 = sp.expand(I2 - I2.subs(x, -1))

        I3 = sp.integrate(I2, x)
        I3 = sp.expand(I3 - I3.subs(x, -1))

        P_coeffs[n] = sympy_poly_to_numpy_coeffs(Pn, x, dtype=dtype)
        I1_coeffs[n] = sympy_poly_to_numpy_coeffs(I1, x, dtype=dtype)
        I2_coeffs[n] = sympy_poly_to_numpy_coeffs(I2, x, dtype=dtype)
        I3_coeffs[n] = sympy_poly_to_numpy_coeffs(I3, x, dtype=dtype)

    return P_coeffs, I1_coeffs, I2_coeffs, I3_coeffs


def pad_coeff_dict_to_common_width(coeff_dict, max_order, common_width, dtype=np.float32):
    mat = np.zeros((max_order + 1, common_width), dtype=dtype)
    for n in range(max_order + 1):
        c = coeff_dict[n]
        mat[n, :len(c)] = c
    return mat


def get_legendre_panel_coefs_sympy_on_custom_panels(order, panel_edges, dtype=np.float32):
    x = sp.symbols("x", real=True)
    P = sp.legendre(order, x)
    Pint = sp.integrate(P, x)

    vals = np.array([float(Pint.subs(x, s)) for s in panel_edges], dtype=np.float64)
    coefs = vals[1:] - vals[:-1]
    coefs *= (2.0 * order + 1.0) / 2.0
    return coefs.astype(dtype)


def differentiate_poly_matrix(poly_mat_np):
    """
    poly_mat_np: shape (n_basis, width), ascending-order coefficients
                 c0 + c1 x + c2 x^2 + ...
    return:
        derivative coefficient matrix with same width
    """
    poly_mat_np = np.asarray(poly_mat_np, dtype=np.float32)
    n_basis, width = poly_mat_np.shape
    out = np.zeros_like(poly_mat_np)

    for k in range(1, width):
        out[:, k - 1] = k * poly_mat_np[:, k]

    return out

# ============================================================
# 2) Legendre anti-derivative approximator
# ============================================================
class LegendreADAFReusable(tf.Module):
    def __init__(
        self,
        xgrid_phys_mapped,
        L=1.0,
        gamma=1.0,
        max_order=10,
        N_p=16,
        init1=0.0,
        init2=0.0,
        init3=0.0,
        dtype=DTYPE,
        seed=0,
        name=None,
    ):
        super().__init__(name=name)

        self.dtype = dtype
        self.L = tf.constant(L, dtype)
        self.gamma = float(gamma)
        self.max_order = int(max_order)
        self.N_p = int(N_p)

        if not (0.0 < self.gamma <= 1.0):
            raise ValueError("gamma must satisfy 0 < gamma <= 1")

        self.x_gamma_np = -1.0 + 2.0 * self.gamma
        self.x_gamma = tf.constant(self.x_gamma_np, dtype=dtype)

        xgrid_phys_mapped = np.asarray(xgrid_phys_mapped, dtype=np.float32)
        if xgrid_phys_mapped.ndim != 1:
            raise ValueError("xgrid_phys_mapped must be 1D")
        if np.any(np.diff(xgrid_phys_mapped) < 0):
            raise ValueError("xgrid_phys_mapped must be monotone increasing")
        if xgrid_phys_mapped[0] < -1.0 - 1e-7 or xgrid_phys_mapped[-1] > self.x_gamma_np + 1e-7:
            raise ValueError("xgrid_phys_mapped must lie in [-1, x_gamma]")

        self.x = tf.constant(xgrid_phys_mapped, dtype=dtype)
        self.xm = self.x + tf.constant(1.0, dtype=dtype)
        self.Nt = int(xgrid_phys_mapped.shape[0])

        if self.gamma == 1.0:
            panel_edges_np = np.linspace(-1.0, 1.0, self.N_p + 1).astype(np.float32)
        else:
            panel_edges_np = np.concatenate(
                (
                    np.linspace(-1.0, self.x_gamma_np, self.N_p - 1, dtype=np.float32)[:-1],
                    np.linspace(self.x_gamma_np, 1.0, 3, dtype=np.float32),
                ),
                axis=0,
            ).astype(np.float32)

        self.panel_edges = tf.constant(panel_edges_np, dtype=dtype)

        rng = np.random.default_rng(seed)
        u0 = rng.uniform(-0.5, 0.5, size=(self.N_p - 1,)).astype(np.float32)
        self.U = tf.Variable(u0, dtype=dtype, name="U")
        

        self.init1 = tf.Variable(float(init1), dtype=dtype, trainable=False, name="init1")
        self.init2 = tf.Variable(float(init2), dtype=dtype, trainable=False, name="init2")
        self.init3 = tf.Variable(float(init3), dtype=dtype, trainable=False, name="init3")

        P_dict, I1_dict, I2_dict, I3_dict = build_legendre_symbolic_tables(
            max_order=self.max_order,
            dtype=np.float32,
        )

        common_width = self.max_order + 4
        self.P_poly_mat = tf.constant(
            pad_coeff_dict_to_common_width(P_dict, self.max_order, common_width), dtype=dtype
        )
        self.I1_poly_mat = tf.constant(
            pad_coeff_dict_to_common_width(I1_dict, self.max_order, common_width), dtype=dtype
        )
        self.I2_poly_mat = tf.constant(
            pad_coeff_dict_to_common_width(I2_dict, self.max_order, common_width), dtype=dtype
        )
        self.I3_poly_mat = tf.constant(
            pad_coeff_dict_to_common_width(I3_dict, self.max_order, common_width), dtype=dtype
        )

        coef_rows = [
            get_legendre_panel_coefs_sympy_on_custom_panels(
                order=n, panel_edges=panel_edges_np, dtype=np.float32
            )
            for n in range(1, self.max_order + 1)
        ]
        self.coef_mat = tf.constant(np.stack(coef_rows, axis=0).astype(np.float32), dtype=dtype)

        self.mean_vec = tf.constant(
            np.ones((self.N_p,), dtype=np.float32) / float(self.N_p), dtype=dtype
        )

        order_idx = tf.constant(np.arange(1, self.max_order + 1), dtype=tf.int32)

        self.P_cache = self._build_basis_cache(self.P_poly_mat, order_idx)
        self.I1_cache = self._build_basis_cache(self.I1_poly_mat, order_idx)
        self.I2_cache = self._build_basis_cache(self.I2_poly_mat, order_idx)
        self.I3_cache = self._build_basis_cache(self.I3_poly_mat, order_idx)

        self.P0_cache = self._eval_poly_matrix(self.P_poly_mat[0:1], self.x)[0]
        self.I1_0_cache = self._eval_poly_matrix(self.I1_poly_mat[0:1], self.x)[0]
        self.I2_0_cache = self._eval_poly_matrix(self.I2_poly_mat[0:1], self.x)[0]
        self.I3_0_cache = self._eval_poly_matrix(self.I3_poly_mat[0:1], self.x)[0]

        P_deriv_np = differentiate_poly_matrix(self.P_poly_mat.numpy())
        self.Pd_poly_mat = tf.constant(P_deriv_np, dtype=dtype)
        
        self.Pd_cache = self._build_basis_cache(self.Pd_poly_mat, order_idx)
        self.Pd0_cache = self._eval_poly_matrix(self.Pd_poly_mat[0:1], self.x)[0]        

    def get_W(self):
        w_last = -tf.reduce_sum(self.U, keepdims=True)
        return tf.concat([self.U, w_last], axis=0)

    def dfdx_from_coeffs(self, a0, A):
        return a0 * self.Pd0_cache + tf.reduce_sum(A[:, None] * self.Pd_cache, axis=0)

    def q_qdot_qddot_qjerk(self, time_scale):
        a0, A = self.coeffs()

        q = self.g2_from_coeffs(a0, A)
        g1 = self.g1_from_coeffs(a0, A)
        f = self.f_from_coeffs(a0, A)
        dfdx = self.dfdx_from_coeffs(a0, A)

        qdot = time_scale * g1
        qddot = (time_scale ** 2) * f
        qjerk = (time_scale ** 3) * dfdx
        return q, qdot, qddot, qjerk

    def _eval_poly_matrix(self, poly_mat, x):
        poly_mat = tf.convert_to_tensor(poly_mat, dtype=self.dtype)
        x = tf.convert_to_tensor(x, dtype=self.dtype)

        n_basis = tf.shape(poly_mat)[0]
        Nt = tf.shape(x)[0]

        y = tf.zeros((n_basis, Nt), dtype=self.dtype)
        coeffs_rev = tf.reverse(poly_mat, axis=[1])

        for c in tf.unstack(coeffs_rev, axis=1):
            y = y * x[None, :] + c[:, None]

        return y

    def _build_basis_cache(self, poly_mat, order_idx):
        sub_mat = tf.gather(poly_mat, order_idx, axis=0)
        return self._eval_poly_matrix(sub_mat, self.x)

    def coeffs(self):
        W = self.get_W()
        a0 = tf.tensordot(self.mean_vec, W, axes=1)
        A = tf.linalg.matvec(self.coef_mat, W)
        return a0, A

    def g1_from_coeffs(self, a0, A):
        return a0 * self.I1_0_cache + tf.reduce_sum(A[:, None] * self.I1_cache, axis=0) + self.init1

    def g2_from_coeffs(self, a0, A):
        return (
            a0 * self.I2_0_cache
            + tf.reduce_sum(A[:, None] * self.I2_cache, axis=0)
            + self.init1 * self.xm
            + self.init2
        )

    def g3_from_coeffs(self, a0, A):
        return (
            a0 * self.I3_0_cache
            + tf.reduce_sum(A[:, None] * self.I3_cache, axis=0)
            + 0.5 * self.init1 * tf.square(self.xm)
            + self.init2 * self.xm
            + self.init3
        )

    def q(self):
        a0, A = self.coeffs()
        return self.g2_from_coeffs(a0, A)

    def f_from_coeffs(self, a0, A):
        return a0 * self.P0_cache + tf.reduce_sum(A[:, None] * self.P_cache, axis=0)

    def q_qdot_qddot(self, time_scale):
        a0, A = self.coeffs()

        q = self.g2_from_coeffs(a0, A)
        g1 = self.g1_from_coeffs(a0, A)
        f = self.f_from_coeffs(a0, A)

        qdot = time_scale * g1
        qddot = (time_scale ** 2) * f
        return q, qdot, qddot


# ============================================================
# 3) UR5e kinematics
# ============================================================
def dh_transform(alpha, a, d, theta):
    theta = tf.cast(theta, DTYPE)
    N = tf.shape(theta)[0]

    ca = tf.constant(np.cos(alpha), dtype=DTYPE)
    sa = tf.constant(np.sin(alpha), dtype=DTYPE)
    ct = tf.cos(theta)
    st = tf.sin(theta)

    row1 = tf.stack([ct, -st * ca, st * sa, tf.constant(a, dtype=DTYPE) * ct], axis=-1)
    row2 = tf.stack([st, ct * ca, -ct * sa, tf.constant(a, dtype=DTYPE) * st], axis=-1)
    row3 = tf.stack(
        [
            tf.zeros_like(theta),
            sa * tf.ones_like(theta),
            ca * tf.ones_like(theta),
            tf.constant(d, dtype=DTYPE) * tf.ones_like(theta),
        ],
        axis=-1,
    )
    row4 = tf.stack(
        [
            tf.zeros_like(theta),
            tf.zeros_like(theta),
            tf.zeros_like(theta),
            tf.ones_like(theta),
        ],
        axis=-1,
    )

    return tf.stack([row1, row2, row3, row4], axis=1)


class UR5eKinematics:
    def __init__(self):
        # UR5e manufacturer DH parameters (Joint 1~6)
        self.alpha = [np.pi / 2.0, 0.0, 0.0, np.pi / 2.0, -np.pi / 2.0, 0.0]
        self.a = [0.0, -0.425, -0.3922, 0.0, 0.0, 0.0]
        self.d = [0.1625, 0.0, 0.0, 0.1333, 0.0997, 0.0996]
        self.theta_offset = [0.0, -np.pi / 2.0, 0.0, -np.pi / 2.0, 0.0, 0.0]

    def forward_all_points(self, q_batch):
        q_batch = tf.cast(q_batch, DTYPE)
        N = tf.shape(q_batch)[0]

        T = tf.eye(4, batch_shape=[N], dtype=DTYPE)
        points = [T[:, :3, 3]]

        for i in range(6):
            theta_i = q_batch[:, i] + tf.constant(self.theta_offset[i], dtype=DTYPE)
            A_i = dh_transform(self.alpha[i], self.a[i], self.d[i], theta_i)
            T = tf.matmul(T, A_i)
            points.append(T[:, :3, 3])

        ee_pos = T[:, :3, 3]
        return points, ee_pos


# ============================================================
# 4) Metrics utilities
# ============================================================
def benchmark_metrics(q, qdot, qddot, qjerk, ee, target_xyz, dt, solve_time):
    int_sq_vel = float(np.sum(np.sum(qdot**2, axis=1)) * dt)
    int_sq_acc = float(np.sum(np.sum(qddot**2, axis=1)) * dt)
    int_sq_jerk = float(np.sum(np.sum(qjerk**2, axis=1)) * dt)

    mean_sq_vel = float(np.mean(np.sum(qdot**2, axis=1)))
    mean_sq_acc = float(np.mean(np.sum(qddot**2, axis=1)))
    mean_sq_jerk = float(np.mean(np.sum(qjerk**2, axis=1)))

    vel_norm_t = np.linalg.norm(qdot, axis=1)
    acc_norm_t = np.linalg.norm(qddot, axis=1)
    jerk_norm_t = np.linalg.norm(qjerk, axis=1)

    metrics = {
        "joint_path_length": float(np.sum(np.linalg.norm(q[1:] - q[:-1], axis=1))),
        "ee_path_length": float(np.sum(np.linalg.norm(ee[1:] - ee[:-1], axis=1))),
        "integrated_squared_velocity": int_sq_vel,
        "integrated_squared_acceleration": int_sq_acc,
        "integrated_squared_jerk": int_sq_jerk,
        "mean_squared_velocity": mean_sq_vel,
        "mean_squared_acceleration": mean_sq_acc,
        "mean_squared_jerk": mean_sq_jerk,
        "max_velocity_norm": float(np.max(vel_norm_t)),
        "max_acceleration_norm": float(np.max(acc_norm_t)),
        "max_jerk_norm": float(np.max(jerk_norm_t)),
        "max_abs_velocity_each_joint": np.max(np.abs(qdot), axis=0),
        "max_abs_acceleration_each_joint": np.max(np.abs(qddot), axis=0),
        "max_abs_jerk_each_joint": np.max(np.abs(qjerk), axis=0),
        "terminal_velocity_norm": float(np.linalg.norm(qdot[-1])),
        "terminal_acceleration_norm": float(np.linalg.norm(qddot[-1])),
        "initial_acceleration_norm": float(np.linalg.norm(qddot[0])),
        "final_position_error": float(np.linalg.norm(ee[-1] - target_xyz)),
        "solve_time": float(solve_time),
    }
    return metrics


def print_metrics_table(name, metrics):
    print(f"\n=== {name} metrics ===")
    keys_scalar = [
        "final_position_error",
        "joint_path_length",
        "ee_path_length",
        "integrated_squared_velocity",
        "integrated_squared_acceleration",
        "integrated_squared_jerk",
        "mean_squared_velocity",
        "mean_squared_acceleration",
        "mean_squared_jerk",
        "max_velocity_norm",
        "max_acceleration_norm",
        "max_jerk_norm",
        "terminal_velocity_norm",
        "terminal_acceleration_norm",
        "initial_acceleration_norm",
        "solve_time",
    ]
    for k in keys_scalar:
        print(f"{k:>32s} : {metrics[k]:.6e}")

    print("\nmax_abs_velocity_each_joint     :", np.array2string(metrics["max_abs_velocity_each_joint"], precision=5))
    print("max_abs_acceleration_each_joint :", np.array2string(metrics["max_abs_acceleration_each_joint"], precision=5))
    print("max_abs_jerk_each_joint         :", np.array2string(metrics["max_abs_jerk_each_joint"], precision=5))


def print_comparison_table(metrics_a, metrics_b, name_a="ADA-L", name_b="Quintic"):
    keys = [
        "final_position_error",
        "joint_path_length",
        "ee_path_length",
        "integrated_squared_velocity",
        "integrated_squared_acceleration",
        "integrated_squared_jerk",
        "max_velocity_norm",
        "max_acceleration_norm",
        "max_jerk_norm",
        "terminal_velocity_norm",
        "terminal_acceleration_norm",
        "initial_acceleration_norm",
        "solve_time",
    ]

    print("\n" + "=" * 98)
    print(f"{'Metric':>32s} | {name_a:>18s} | {name_b:>18s}")
    print("=" * 98)
    for k in keys:
        print(f"{k:>32s} | {metrics_a[k]:18.6e} | {metrics_b[k]:18.6e}")
    print("=" * 98)


# ============================================================
# 4b) Thin proxy for plot helpers (self.models[j].U / .panel_edges)
# ============================================================
class _RowSliceVar:
    def __init__(self, parent, row):
        self._p = parent; self._r = row
    @property
    def shape(self): return self._p.shape[1:]
    def numpy(self): return self._p[self._r].numpy()
    def assign(self, v):
        idx = tf.constant([[self._r]], dtype=tf.int32)
        self._p.assign(tf.tensor_scatter_nd_update(
            self._p, idx, tf.reshape(tf.cast(v, self._p.dtype), (1, -1))))

class _JointProxy:
    def __init__(self, planner, j, panel_edges_tf):
        self._planner = planner; self._j = j
        self.panel_edges = panel_edges_tf
    @property
    def U(self): return _RowSliceVar(self._planner.U_batch, self._j)

def _make_joint_proxies(planner, panel_edges_np):
    pe = tf.constant(panel_edges_np, dtype=DTYPE)
    return [_JointProxy(planner, j, pe) for j in range(6)]


# ============================================================
# 5) ADA-L planner
# ============================================================
class ADALTrajectoryPlanner:
    def __init__(
        self,
        q0,
        target_xyz,
        T_final=1.0,
        Nt=201,
        gamma=1.0,
        max_order=10,
        N_p=16,
        seed=0,
    ):
        self.q0 = np.asarray(q0, dtype=np.float32).reshape(6,)
        self.target_xyz = tf.constant(np.asarray(target_xyz, dtype=np.float32).reshape(3,), dtype=DTYPE)
        self.T_final = float(T_final)
        self.Nt = int(Nt)
        self.dt = self.T_final / (self.Nt - 1)
        self.gamma = float(gamma)
        self.N_p = int(N_p)
        self.max_order = int(max_order)
        self.optimizer = None

        self.t_np = np.linspace(0.0, self.T_final, self.Nt).astype(np.float32)
        self.xgrid_np = (-1.0 + 2.0 * gamma * (self.t_np / self.T_final)).astype(np.float32)
        self.robot = UR5eKinematics()

        # ---- shared basis matrices ----
        P_dict, I1_dict, I2_dict, I3_dict = build_legendre_symbolic_tables(max_order=max_order)
        cw = max_order + 4
        P_np  = pad_coeff_dict_to_common_width(P_dict,  max_order, cw)
        I1_np = pad_coeff_dict_to_common_width(I1_dict, max_order, cw)
        I2_np = pad_coeff_dict_to_common_width(I2_dict, max_order, cw)
        x_tf  = tf.constant(self.xgrid_np, dtype=DTYPE)
        oi    = np.arange(1, max_order + 1)

        def _eval(mat_np, x):
            mat = tf.constant(mat_np, dtype=DTYPE)
            y = tf.zeros((mat_np.shape[0], x.shape[0]), dtype=DTYPE)
            for c in tf.unstack(tf.reverse(mat, axis=[1]), axis=1):
                y = y * x[None, :] + c[:, None]
            return y

        self._P_cache    = _eval(P_np[oi],   x_tf)
        self._I1_cache   = _eval(I1_np[oi],  x_tf)
        self._I2_cache   = _eval(I2_np[oi],  x_tf)
        self._P0_cache   = _eval(P_np[0:1],  x_tf)[0]
        self._I1_0_cache = _eval(I1_np[0:1], x_tf)[0]
        self._I2_0_cache = _eval(I2_np[0:1], x_tf)[0]
        Pd_np = differentiate_poly_matrix(P_np)
        self._Pd_cache   = _eval(Pd_np[oi],  x_tf)
        self._Pd0_cache  = _eval(Pd_np[0:1], x_tf)[0]

        if gamma == 1.0:
            panel_edges_np = np.linspace(-1.0, 1.0, N_p + 1).astype(np.float32)
        else:
            xg = -1.0 + 2.0 * gamma
            panel_edges_np = np.concatenate((
                np.linspace(-1.0, xg, N_p - 1, dtype=np.float32)[:-1],
                np.linspace(xg, 1.0, 3, dtype=np.float32),
            )).astype(np.float32)

        coef_rows = [get_legendre_panel_coefs_sympy_on_custom_panels(n, panel_edges_np)
                     for n in range(1, max_order + 1)]
        self._coef_mat = tf.constant(np.stack(coef_rows).astype(np.float32), dtype=DTYPE)
        self._mean_vec = tf.constant(np.ones((N_p,), dtype=np.float32) / float(N_p), dtype=DTYPE)
        self._init2    = tf.constant(self.q0, dtype=DTYPE)

        rng = np.random.default_rng(seed)
        U0 = np.stack([rng.uniform(-0.5, 0.5, size=(N_p-1,)).astype(np.float32)
                       for _ in range(6)], axis=0)
        self.U_batch = tf.Variable(U0, dtype=DTYPE, name="U_batch")
        self.models  = _make_joint_proxies(self, panel_edges_np)

        # UR5e joint position limits: elbow (joint 3) is +/-pi, others +/-2*pi
        self.q_min = tf.constant([-2*np.pi, -2*np.pi, -np.pi, -2*np.pi, -2*np.pi, -2*np.pi], dtype=DTYPE)
        self.q_max = tf.constant([ 2*np.pi,  2*np.pi,  np.pi,  2*np.pi,  2*np.pi,  2*np.pi], dtype=DTYPE)
        # UR5e velocity limit: 180 deg/s = pi rad/s
        self.qdot_max  = tf.constant(np.pi, dtype=DTYPE)
        # UR5e acceleration limit: 2292 deg/s^2 = 40 rad/s^2
        self.qddot_max = tf.constant(40.0, dtype=DTYPE)

        self.history = {
            "loss": [], "goal": [], "path": [], "acc": [], "jerk": [],
            "limit": [], "terminal_vel": [], "terminal_acc": [], "initial_acc": [],
        }

    def get_trainable_variables(self):
        return [self.U_batch]

    def _batch_q_qdot_qddot_qjerk(self, time_scale):
        w_last = -tf.reduce_sum(self.U_batch, axis=1, keepdims=True)
        W  = tf.concat([self.U_batch, w_last], axis=1)
        a0 = tf.linalg.matvec(W, self._mean_vec)
        A  = tf.matmul(W, tf.transpose(self._coef_mat))
        def _ap(a0_, A_, c0_, cache_):
            return a0_[:, None] * c0_[None, :] + tf.reduce_sum(
                A_[:, :, None] * cache_[None, :, :], axis=1)
        g2 = _ap(a0, A, self._I2_0_cache, self._I2_cache) + self._init2[:, None]
        g1 = _ap(a0, A, self._I1_0_cache, self._I1_cache)
        f  = _ap(a0, A, self._P0_cache,   self._P_cache)
        df = _ap(a0, A, self._Pd0_cache,  self._Pd_cache)
        return (tf.transpose(g2), tf.transpose(time_scale * g1),
                tf.transpose((time_scale**2) * f), tf.transpose((time_scale**3) * df))

    def q_traj(self):
        c = tf.constant(2.0 * self.gamma / self.T_final, dtype=DTYPE)
        return self._batch_q_qdot_qddot_qjerk(c)[0]

    def q_traj_with_derivatives(self):
        c = tf.constant(2.0 * self.gamma / self.T_final, dtype=DTYPE)
        q, qdot, qddot, _ = self._batch_q_qdot_qddot_qjerk(c)
        return q, qdot, qddot

    def q_traj_with_derivatives_and_jerk(self):
        c = tf.constant(2.0 * self.gamma / self.T_final, dtype=DTYPE)
        return self._batch_q_qdot_qddot_qjerk(c)

    def finite_diff(self, y):
        return (y[1:] - y[:-1]) / tf.constant(self.dt, dtype=DTYPE)

    def terminal_ee_from_q(self, q):
        qT = q[-1:,:]                 # shape (1,6)
        _, eeT = self.robot.forward_all_points(qT)
        return eeT[0]                 # shape (3,)

    def full_ee_from_q(self, q):
        _, ee = self.robot.forward_all_points(q)
        return ee

    def compute_loss_train(self):
        q, qdot, qddot, qjerk = self.q_traj_with_derivatives_and_jerk()

        eeT = self.terminal_ee_from_q(q)

        goal_loss = tf.reduce_mean(tf.square(eeT - self.target_xyz))
        path_loss = tf.reduce_mean(tf.square(qdot))
        acc_loss = tf.reduce_mean(tf.square(qddot))
        jerk_loss = tf.reduce_mean(tf.square(qjerk))

        # L1 penalty with 90% safety margin (Ex1 pattern).
        margin_frac = tf.constant(0.9, dtype=DTYPE)
        q_mid  = 0.5 * (self.q_min + self.q_max)
        q_half = 0.5 * (self.q_max - self.q_min)
        q_lo_safe = (q_mid - margin_frac * q_half)[None, :]
        q_hi_safe = (q_mid + margin_frac * q_half)[None, :]
        lv = tf.nn.relu(q_lo_safe - q)
        uv = tf.nn.relu(q - q_hi_safe)
        limit_loss = tf.reduce_mean(lv + uv)
        vel_violation = tf.nn.relu(tf.abs(qdot)  - margin_frac * self.qdot_max)
        acc_violation = tf.nn.relu(tf.abs(qddot) - margin_frac * self.qddot_max)
        vel_limit_loss = tf.reduce_mean(vel_violation)
        acc_limit_loss = tf.reduce_mean(acc_violation)

        terminal_vel_loss = tf.reduce_mean(tf.square(qdot[-1]))
        terminal_acc_loss = tf.reduce_mean(tf.square(qddot[-1]))
        initial_acc_loss = tf.reduce_mean(tf.square(qddot[0]))

        loss = (
            5000000.0 * goal_loss
            + 2000.0 * path_loss
            # + 1.0 * acc_loss
            + 1.0 * jerk_loss
            + 1000.0 * limit_loss
            + 1000.0 * vel_limit_loss
            + 1000.0 * acc_limit_loss
            #+ 3.0 * terminal_vel_loss
            + 500.0 * initial_acc_loss
            + 500.0 * terminal_acc_loss
        )

        return {
            "loss": loss,
            "goal": goal_loss,
            "path": path_loss,
            "acc": acc_loss,
            "jerk": jerk_loss,
            "limit": limit_loss,
            "vel_limit": vel_limit_loss,
            "acc_limit": acc_limit_loss,
            "terminal_vel": terminal_vel_loss,
            "terminal_acc": terminal_acc_loss,
            "initial_acc": initial_acc_loss,
            "q": q,
            "qdot": qdot,
            "qddot": qddot,
            "qjerk": qjerk,
            "eeT": eeT,
        }

    def pack_weights_np(self):
        return self.U_batch.numpy().reshape(-1).astype(np.float64)

    def unpack_weights_np(self, w_flat):
        self.U_batch.assign(tf.constant(
            np.asarray(w_flat, dtype=np.float32).reshape(self.U_batch.shape), dtype=DTYPE))

    @tf.function(reduce_retracing=True)
    def _loss_and_grad_tf(self):
        with tf.GradientTape() as tape:
            loss = self.compute_loss_train()["loss"]
        grads = tape.gradient(loss, self.get_trainable_variables())
        return loss, grads[0]

    def loss_and_grad_np(self, w_flat):
        self.unpack_weights_np(w_flat)
        loss_tf, grad_tf = self._loss_and_grad_tf()
        grad_flat = grad_tf.numpy().reshape(-1).astype(np.float64)
        return float(loss_tf.numpy()), grad_flat

    def run_lbfgs(self, maxiter=2000):
        w0 = self.pack_weights_np()
        t0 = time.time()
        res = scipy.optimize.minimize(
            fun=self.loss_and_grad_np,
            x0=w0,
            jac=True,
            method="L-BFGS-B",
            options={
                "maxiter": int(maxiter),
                "maxfun": 50000,
                "maxcor": 50,
                "maxls": 50,
                "ftol": 1e-10,
                "gtol": 1e-10,
                "iprint": -1,
            },
        )
        elapsed = time.time() - t0
        self.unpack_weights_np(res.x)

        out = self.compute_loss_train()
        print(
            f"[ADA-L-LBFGS] done in {elapsed:.2f} sec | "
            f"loss={out['loss'].numpy():.6e} | "
            f"goal={out['goal'].numpy():.6e} | "
            f"tvel={out['terminal_vel'].numpy():.6e} | "
            f"tacc={out['terminal_acc'].numpy():.6e} | "
            f"msg={str(res.message).splitlines()[0]}"
        )
        return res, elapsed

    @tf.function
    def train_step(self):
        vars_ = self.get_trainable_variables()

        with tf.GradientTape() as tape:
            out = self.compute_loss_train()
            loss = out["loss"]

        grads = tape.gradient(loss, vars_)
        grads_and_vars = [(g, v) for g, v in zip(grads, vars_) if g is not None]

        if len(grads_and_vars) > 0:
            self.optimizer.apply_gradients(grads_and_vars)
        else:
            tf.print("WARNING: all gradients are None in ADA-L train_step")

        return (
            loss,
            out["goal"],
            out["path"],
            out["acc"],
            out["jerk"],
            out["limit"],
            out["terminal_vel"],
            out["terminal_acc"],
            out["initial_acc"],
        )


    def train(self, epochs=1000, lr=3e-2, print_every=200, use_lbfgs=True, lbfgs_maxiter=2000):
        self.optimizer = tf.keras.optimizers.Adam(learning_rate=lr)

        t0 = time.time()
        for ep in range(1, epochs + 1):
            loss, goal, path, acc, jerk, limit, terminal_vel, terminal_acc, initial_acc = self.train_step()

            self.history["loss"].append(float(loss.numpy()))
            self.history["goal"].append(float(goal.numpy()))
            self.history["path"].append(float(path.numpy()))
            self.history["acc"].append(float(acc.numpy()))
            self.history["jerk"].append(float(jerk.numpy()))
            self.history["limit"].append(float(limit.numpy()))
            self.history["terminal_vel"].append(float(terminal_vel.numpy()))
            self.history["terminal_acc"].append(float(terminal_acc.numpy()))
            self.history["initial_acc"].append(float(initial_acc.numpy()))

            if ep % print_every == 0 or ep == 1:
                print(
                    f"[ADA-L-Adam {ep:5d}/{epochs}] "
                    f"loss={loss.numpy():.6e} | "
                    f"goal={goal.numpy():.6e} | "
                    f"jerk={jerk.numpy():.6e} | "
                    f"tvel={terminal_vel.numpy():.6e} | "
                    f"tacc={terminal_acc.numpy():.6e} | "
                    f"iacc={initial_acc.numpy():.6e}"
                )

        adam_elapsed = time.time() - t0
        print(f"[ADA-L] Adam finished in {adam_elapsed:.2f} sec")

        lbfgs_elapsed = 0.0
        if use_lbfgs:
            _, lbfgs_elapsed = self.run_lbfgs(maxiter=lbfgs_maxiter)

        return adam_elapsed + lbfgs_elapsed

    def results(self):
        q, qdot, qddot, qjerk = self.q_traj_with_derivatives_and_jerk()
        ee = self.full_ee_from_q(q)

        return {
            "q": q.numpy(),
            "qdot": qdot.numpy(),
            "qddot": qddot.numpy(),
            "qjerk": qjerk.numpy(),
            "ee": ee.numpy(),
            "t": self.t_np.copy(),
            "target_xyz": self.target_xyz.numpy(),
        }


# ============================================================
# 6) Terminal IK solver for quintic baseline
# ============================================================
class TerminalIKSolver:
    def __init__(self, robot, q0, target_xyz, q_min=None, q_max=None):
        self.robot = robot
        self.q0 = tf.constant(np.asarray(q0, dtype=np.float32).reshape(6,), dtype=DTYPE)
        self.target_xyz = tf.constant(np.asarray(target_xyz, dtype=np.float32).reshape(3,), dtype=DTYPE)

        if q_min is None:
            q_min = [-2 * np.pi] * 6
        if q_max is None:
            q_max = [2 * np.pi] * 6

        self.q_min = tf.constant(np.asarray(q_min, dtype=np.float32), dtype=DTYPE)
        self.q_max = tf.constant(np.asarray(q_max, dtype=np.float32), dtype=DTYPE)

        self.q_var = tf.Variable(self.q0.numpy(), dtype=DTYPE, trainable=True, name="terminal_ik_q")
        self.optimizer = None

        self.history = {
            "loss": [],
            "goal": [],
            "reg": [],
            "limit": [],
            "pos_err": [],
        }

    @tf.function
    def step(self):
        with tf.GradientTape() as tape:
            _, ee = self.robot.forward_all_points(self.q_var[None, :])
            ee = ee[0]

            goal_loss = tf.reduce_mean(tf.square(ee - self.target_xyz))
            reg_loss = 1e-3 * tf.reduce_mean(tf.square(self.q_var - self.q0))

            lower_violation = tf.nn.relu(self.q_min - self.q_var)
            upper_violation = tf.nn.relu(self.q_var - self.q_max)
            limit_loss = tf.reduce_mean(tf.square(lower_violation) + tf.square(upper_violation))

            loss = 20000.0 * goal_loss + reg_loss + 20.0 * limit_loss

        grad = tape.gradient(loss, self.q_var)

        # Guard against the abnormal case where gradient is None
        if grad is not None:
            self.optimizer.apply_gradients([(grad, self.q_var)])
        else:
            tf.print("WARNING: IK gradient is None")

        pos_err = tf.norm(ee - self.target_xyz)

        return loss, goal_loss, reg_loss, limit_loss, pos_err

    def solve(self, epochs=3000, lr=3e-2, print_every=300):
        # Reset initial value on every solve
        self.q_var.assign(self.q0)
        self.optimizer = tf.keras.optimizers.Adam(learning_rate=lr)

        # history reset
        for k in self.history:
            self.history[k] = []

        t0 = time.time()
        for ep in range(1, epochs + 1):
            loss, goal_loss, reg_loss, limit_loss, pos_err = self.step()

            self.history["loss"].append(float(loss.numpy()))
            self.history["goal"].append(float(goal_loss.numpy()))
            self.history["reg"].append(float(reg_loss.numpy()))
            self.history["limit"].append(float(limit_loss.numpy()))
            self.history["pos_err"].append(float(pos_err.numpy()))

            if ep % print_every == 0 or ep == 1:
                print(
                    f"[IK {ep:5d}/{epochs}] "
                    f"loss={loss.numpy():.6e} | "
                    f"pos_err={pos_err.numpy():.6e}"
                )

        elapsed = time.time() - t0
        _, ee = self.robot.forward_all_points(self.q_var[None, :])
        ee = ee[0].numpy()

        return self.q_var.numpy(), ee, elapsed, self.history


# ============================================================
# 7) Quintic polynomial baseline
# ============================================================
class JointQuinticBaseline:
    def __init__(self, robot, q0, target_xyz, T_final=1.0, Nt=201):
        self.robot = robot
        self.q0 = np.asarray(q0, dtype=np.float32).reshape(6,)
        self.target_xyz = np.asarray(target_xyz, dtype=np.float32).reshape(3,)
        self.T_final = float(T_final)
        self.Nt = int(Nt)
        self.dt = self.T_final / (self.Nt - 1)
        self.t = np.linspace(0.0, self.T_final, self.Nt).astype(np.float32)

    def solve_terminal_q(self, ik_epochs=3000, ik_lr=3e-2):
        solver = TerminalIKSolver(self.robot, self.q0, self.target_xyz)
        qT, eeT, ik_time, ik_history = solver.solve(
            epochs=ik_epochs,
            lr=ik_lr,
            print_every=max(1, ik_epochs // 10),
        )
        pos_err = np.linalg.norm(eeT - self.target_xyz)
        print(f"[Quintic-IK] done in {ik_time:.2f} sec | final_pos_err={pos_err:.6e}")
        return qT, ik_time, ik_history

    def generate_quintic(self, qT):
        t = self.t[:, None]
        T = self.T_final
        s = t / T

        h = 10.0 * s**3 - 15.0 * s**4 + 6.0 * s**5
        dh_dt = (30.0 * s**2 - 60.0 * s**3 + 30.0 * s**4) / T
        d2h_dt2 = (60.0 * s - 180.0 * s**2 + 120.0 * s**3) / (T**2)
        d3h_dt3 = (60.0 - 360.0 * s + 360.0 * s**2) / (T**3)

        dq = (qT - self.q0)[None, :]

        q = self.q0[None, :] + h * dq
        qdot = dh_dt * dq
        qddot = d2h_dt2 * dq
        qjerk = d3h_dt3 * dq

        q = q.astype(np.float32)
        qdot = qdot.astype(np.float32)
        qddot = qddot.astype(np.float32)
        qjerk = qjerk.astype(np.float32)

        _, ee = self.robot.forward_all_points(tf.constant(q, dtype=DTYPE))
        ee = ee.numpy().astype(np.float32)

        return q, qdot, qddot, qjerk, ee

    def run(self, ik_epochs=3000, ik_lr=3e-2):
        t0 = time.time()
        qT, ik_time, ik_history = self.solve_terminal_q(ik_epochs=ik_epochs, ik_lr=ik_lr)
        q, qdot, qddot, qjerk, ee = self.generate_quintic(qT)
        total_time = time.time() - t0

        return {
            "qT": qT,
            "q": q,
            "qdot": qdot,
            "qddot": qddot,
            "qjerk": qjerk,
            "ee": ee,
            "t": self.t.copy(),
            "ik_time": ik_time,
            "ik_history": ik_history,
            "solve_time": total_time,
        }

# ============================================================
# 8) Visualization
# ============================================================
def smart_tick_formatter(x, pos):
    if np.isclose(x, round(x)):
        return f"{int(round(x))}"
    return f"{x:.2f}".rstrip("0").rstrip(".")


def apply_tick_style_2d(ax):
    ax.xaxis.set_major_locator(MaxNLocator(nbins=5))
    ax.yaxis.set_major_locator(MaxNLocator(nbins=5))
    formatter = FuncFormatter(smart_tick_formatter)
    ax.xaxis.set_major_formatter(formatter)
    ax.yaxis.set_major_formatter(formatter)


def apply_tick_style_3d(ax):
    # ax.xaxis.set_major_locator(MaxNLocator(nbins=4))
    # ax.yaxis.set_major_locator(MaxNLocator(nbins=4))
    # ax.zaxis.set_major_locator(MaxNLocator(nbins=4))
    # ax.xaxis.set_major_formatter(FormatStrFormatter("%.1f"))
    # ax.yaxis.set_major_formatter(FormatStrFormatter("%.1f"))
    # ax.zaxis.set_major_formatter(FormatStrFormatter("%.1f"))

    ax.set_xticks(np.arange(-1.0, 2.0, 0.25))
    ax.set_yticks(np.arange(-1.0, 2.0, 0.25))
    ax.set_zticks(np.arange(-1.0, 2.0, 0.25))

    formatter = FuncFormatter(smart_tick_formatter)
    ax.xaxis.set_major_formatter(formatter)
    ax.yaxis.set_major_formatter(formatter)
    ax.zaxis.set_major_formatter(formatter)


def style_2d_axes(ax):
    ax.grid(False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(True)
    ax.spines["bottom"].set_visible(True)
    apply_tick_style_2d(ax)


def style_3d_axes(ax):
    ax.grid(True)
    ax.view_init(elev=15, azim=-135)
    ax.tick_params(axis="x", pad=0.2)
    ax.tick_params(axis="y", pad=0.2)
    ax.tick_params(axis="z", pad=3)
    apply_tick_style_3d(ax)


def set_equal_3d_axes(ax, xyz):
    xyz = np.asarray(xyz, dtype=np.float32).reshape(-1, 3)
    mins = xyz.min(axis=0)
    maxs = xyz.max(axis=0)
    centers = 0.5 * (mins + maxs)
    L = float(np.max(maxs - mins))
    if L < 1e-9:
        L = 0.1

    ax.set_xlim(centers[0] - L / 2, centers[0] + L / 2)
    ax.set_ylim(centers[1] - L / 2, centers[1] + L / 2)
    ax.set_zlim(centers[2] - L / 2, centers[2] + L / 2)
    ax.set_box_aspect([1, 1, 1])
    apply_tick_style_3d(ax)


def plot_joint_bundle_single(t, y, title, y_label, labels, save_path=None, show=True):
    fig, ax = plt.subplots(figsize=(5.2, 2.5))
    for j in range(y.shape[1]):
        ax.plot(t, y[:, j], linewidth=1.0, label=labels[j])

    ax.set_xlabel(r"Time, $t$ (s)")
    ax.set_ylabel(y_label)
    # ax.set_title(title)
    style_2d_axes(ax)
    fig.tight_layout()

    if save_path is not None:
        fig.savefig(save_path, dpi=300, bbox_inches="tight")
    if show:
        plt.show()
    else:
        plt.close(fig)


from matplotlib.lines import Line2D

def plot_joint_trajectories_compare(t, q_a, q_b, label_a="ADA-L", label_b="Quintic", save_path=None, show=True, q0=None):
    joint_labels = [f"Joint {j+1}" for j in range(6)]

    fig, axes = plt.subplots(3, 2, figsize=(5.6, 6.8), sharex=True)
    axes = axes.ravel()

    for j, ax in enumerate(axes):
        ax.plot(t, q_a[:, j], linewidth=1.0, label=label_a if j == 0 else None)
        ax.plot(t, q_b[:, j], "--", linewidth=1.0, label=label_b if j == 0 else None)

        ax.set_xlabel(r"Time, $t$ (s)")
        ax.set_ylabel(r"Angle, $q$ (rad)")
        style_2d_axes(ax)
        if j == 5 and q0 is not None and np.abs(np.max(q_a[:, j]) - np.min(q_a[:, j])) < 0.1:
            c = float(q0[5])
            ax.set_ylim(c - 0.1, c + 0.1)
            ax.set_yticks([c - 0.1, c, c + 0.1])

        # Dummy artist with no symbol
        dummy = Line2D([], [], linestyle='None', marker=None, linewidth=0, color='none')

        # Show only the joint name at the best location inside the subplot
        leg_joint = ax.legend(
            [dummy],
            [joint_labels[j]],
            loc="best",
            frameon=False,
            handlelength=0,
            handletextpad=0,
            borderpad=0.2,
            fontsize=10
        )
        ax.add_artist(leg_joint)

    fig.tight_layout()

    if save_path is not None:
        fig.savefig(save_path, dpi=300, bbox_inches="tight")
    if show:
        plt.show()
    else:
        plt.close(fig)


def plot_norm_profiles_compare(t, y_a, y_b, title, y_label, label_a="ADA-L", label_b="Quintic", save_path=None, show=True):
    n = min(len(t), len(y_a), len(y_b))
    fig, ax = plt.subplots(figsize=(5.2, 2.5))
    ax.plot(t[:n], y_a[:n], linewidth=1.0, label=label_a)
    ax.plot(t[:n], y_b[:n], "--", linewidth=1.0, label=label_b)

    ax.set_xlabel(r"Time, $t$ (s)")
    ax.set_ylabel(y_label)
    # ax.set_title(title)
    style_2d_axes(ax)
    fig.tight_layout()

    if save_path is not None:
        fig.savefig(save_path, dpi=300, bbox_inches="tight")
    if show:
        plt.show()
    else:
        plt.close(fig)


def plot_ee_paths_compare(ee_a, ee_b, target_xyz, label_a="ADA-L", label_b="Quintic", save_path=None, show=True, xy_center=None):
    fig = plt.figure(figsize=(6.6, 5.6))
    ax = fig.add_subplot(111, projection="3d")
    ax.plot(ee_a[:, 0], ee_a[:, 1], ee_a[:, 2], linewidth=2.2, label=label_a)
    ax.plot(ee_b[:, 0], ee_b[:, 1], ee_b[:, 2], "--", linewidth=2.2, label=label_b)
    ax.scatter([target_xyz[0]], [target_xyz[1]], [target_xyz[2]], s=110, marker="*", color="k", label="target")
    ax.scatter([ee_a[0, 0]], [ee_a[0, 1]], [ee_a[0, 2]], s=42, color="k", label="start")

    style_3d_axes(ax)
    ax.set_xlabel(r"$x$ (m)", labelpad=2)
    ax.set_ylabel(r"$y$ (m)", labelpad=2)
    ax.set_zlabel(r"$z$ (m)", labelpad=2)
    ax.set_title(f"End-effector path: {label_a} vs. {label_b}")
    # set_equal_3d_axes(ax, np.vstack([ee_a, ee_b, np.asarray(target_xyz).reshape(1, 3)]))

    if xy_center is not None:
        cx, cy = float(xy_center[0]), float(xy_center[1])
        ax.set_xlim(cx - 0.5, cx + 0.5)
        ax.set_ylim(cy - 0.5, cy + 0.5)
    else:
        ax.set_xlim(-0.75, 0.25)
        ax.set_ylim(-0.5, 0.5)
    ax.set_zlim(0, 1.0)
    # ax.set_box_aspect((0.70, 0.80, 0.85))

    ax.legend(loc="upper center", ncol=2, frameon=True)
    fig.tight_layout()

    if save_path is not None:
        fig.savefig(save_path, dpi=300)
    if show:
        plt.show()
    else:
        plt.close(fig)


def plot_robot_snapshots_generic(robot, q, target_xyz, time_grid, n_show=6, title="Robot snapshots", save_path=None, show=True, xy_center=None):
    idxs = np.linspace(0, q.shape[0] - 1, n_show).astype(int)

    L1 = 0.11336
    time_grid = np.asarray(time_grid, dtype=np.float32).reshape(-1)
    if len(time_grid) != q.shape[0]:
        raise ValueError("time_grid length must match q.shape[0]")

    norm = plt.Normalize(vmin=float(time_grid.min()), vmax=float(time_grid.max()))
    cmap = plt.cm.coolwarm

    fig = plt.figure(figsize=(6.6, 5.6))
    ax = fig.add_subplot(111, projection="3d")
    all_xyz = []

    for idx in idxs:
        qk = tf.constant(q[idx:idx+1], dtype=DTYPE)
        color = cmap(norm(float(time_grid[idx])))

        points, _ = robot.forward_all_points(qk)
        pts_np = [p.numpy()[0] for p in points]

        T = tf.eye(4, batch_shape=[1], dtype=DTYPE)
        T01 = None
        for i in range(6):
            theta_i = qk[:, i] + tf.constant(robot.theta_offset[i], dtype=DTYPE)
            A_i = dh_transform(robot.alpha[i], robot.a[i], robot.d[i], theta_i)
            T = tf.matmul(T, A_i)
            if i == 0:
                T01 = T

        o1 = pts_np[1]
        o2 = pts_np[2]
        z1_dir = T01[:, :3, 2].numpy()[0]
        p1_base = o1 + L1 * z1_dir
        p2_base = p1_base + (o2 - o1)

        pts_line = np.asarray(pts_np[:2] + [p1_base, p2_base] + pts_np[2:])
        pts_scatter = np.asarray(pts_np)

        ax.plot(pts_line[:, 0], pts_line[:, 1], pts_line[:, 2], "-", color=color, linewidth=2.0, alpha=0.95)
        ax.scatter(pts_scatter[:, 0], pts_scatter[:, 1], pts_scatter[:, 2], color=[color], s=26, depthshade=True)
        all_xyz.append(pts_line)

    ax.scatter([target_xyz[0]], [target_xyz[1]], [target_xyz[2]], s=120, marker="*", color="k", label="target")
    ax.scatter([all_xyz[0][-1, 0]], [all_xyz[0][-1, 1]], [all_xyz[0][-1, 2]], s=100, color="k", label="start")

    style_3d_axes(ax)
    ax.set_xlabel(r"$x$ (m)", labelpad=2)
    ax.set_ylabel(r"$y$ (m)", labelpad=2)
    ax.set_zlabel(r"$z$ (m)", labelpad=2)
    ax.set_title(title)

    sm = plt.cm.ScalarMappable(cmap=cmap, norm=norm)
    sm.set_array([])
    cbar = fig.colorbar(sm, ax=ax, pad=0.01, shrink=0.75)
    cbar.set_label(r"Time, $t$ (s)")
    cbar.ax.yaxis.set_major_locator(MaxNLocator(nbins=4))
    cbar.ax.yaxis.set_major_formatter(FormatStrFormatter("%.1f"))

    # set_equal_3d_axes(ax, np.vstack(all_xyz + [np.asarray(target_xyz).reshape(1, 3)]))
    if xy_center is not None:
        cx, cy = float(xy_center[0]), float(xy_center[1])
        ax.set_xlim(cx - 0.5, cx + 0.5)
        ax.set_ylim(cy - 0.5, cy + 0.5)
    else:
        ax.set_xlim(-0.75, 0.25)
        ax.set_ylim(-0.5, 0.5)
    ax.set_zlim(0, 1.0)
    # ax.set_box_aspect((0.70, 0.80, 0.85))

    ax.legend(frameon=True)
    fig.tight_layout()
    # fig.subplots_adjust(left=0., right=0.90, top=0.90, bottom=0.08)

    if save_path is not None:
        fig.savefig(save_path, dpi=300)
    if show:
        plt.show()
    else:
        plt.close(fig)

# animate_robot_trajectory
# ============================================================
def animate_robot_trajectory(robot, q, target_xyz, time_grid, save_path="robot_animation.mp4", fps=30, title="Robot trajectory", xy_center=None):
    """
    Render the full trajectory frame-by-frame and save as mp4.
    Uses the same visualization style as plot_robot_snapshots_generic.

    Parameters
    ----------
    robot       : UR5eKinematics (must have forward_all_points method)
    q           : np.ndarray, shape (N, 6)
    target_xyz  : array-like, shape (3,)
    time_grid   : array-like, shape (N,)
    save_path   : str  output mp4 path
    fps         : int  frames per second
    title       : str  axes title
    """
    from matplotlib.animation import FuncAnimation, FFMpegWriter
    from matplotlib.ticker import MaxNLocator, FormatStrFormatter

    L1 = 0.11336
    time_grid = np.asarray(time_grid, dtype=np.float32).reshape(-1)
    N = q.shape[0]
    if len(time_grid) != N:
        raise ValueError("time_grid length must match q.shape[0]")

    norm = plt.Normalize(vmin=float(time_grid.min()), vmax=float(time_grid.max()))
    cmap = plt.cm.coolwarm

    # Pre-compute all poses
    all_pts_line = []
    all_pts_scatter = []
    for k in range(N):
        qk = tf.constant(q[k:k+1], dtype=DTYPE)
        points, _ = robot.forward_all_points(qk)
        pts_np = [p.numpy()[0] for p in points]

        T = tf.eye(4, batch_shape=[1], dtype=DTYPE)
        T01 = None
        for i in range(6):
            theta_i = qk[:, i] + tf.constant(robot.theta_offset[i], dtype=DTYPE)
            A_i = dh_transform(robot.alpha[i], robot.a[i], robot.d[i], theta_i)
            T = tf.matmul(T, A_i)
            if i == 0:
                T01 = T

        o1 = pts_np[1]
        o2 = pts_np[2]
        z1_dir = T01[:, :3, 2].numpy()[0]
        p1_base = o1 + L1 * z1_dir
        p2_base = p1_base + (o2 - o1)

        all_pts_line.append(np.asarray(pts_np[:2] + [p1_base, p2_base] + pts_np[2:]))
        all_pts_scatter.append(np.asarray(pts_np))

    fig = plt.figure(figsize=(6.6, 5.6))
    ax = fig.add_subplot(111, projection="3d")

    # Fixed elements
    ax.scatter([target_xyz[0]], [target_xyz[1]], [target_xyz[2]],
               s=120, marker="*", color="k", label="target", zorder=5)
    ax.scatter([all_pts_line[0][-1, 0]], [all_pts_line[0][-1, 1]], [all_pts_line[0][-1, 2]],
               s=100, color="k", label="start", zorder=5)

    style_3d_axes(ax)
    ax.set_xlabel(r"$x$ (m)", labelpad=2)
    ax.set_ylabel(r"$y$ (m)", labelpad=2)
    ax.set_zlabel(r"$z$ (m)", labelpad=2)
    ax.set_title(title)
    if xy_center is not None:
        cx, cy = float(xy_center[0]), float(xy_center[1])
        ax.set_xlim(cx - 0.5, cx + 0.5)
        ax.set_ylim(cy - 0.5, cy + 0.5)
    else:
        ax.set_xlim(-0.75, 0.25)
        ax.set_ylim(-0.5, 0.5)
    ax.set_zlim(0, 1.0)
    ax.legend(frameon=True)

    sm = plt.cm.ScalarMappable(cmap=cmap, norm=norm)
    sm.set_array([])
    cbar = fig.colorbar(sm, ax=ax, pad=0.01, shrink=0.75)
    cbar.set_label(r"Time, $t$ (s)")
    cbar.ax.yaxis.set_major_locator(MaxNLocator(nbins=4))
    cbar.ax.yaxis.set_major_formatter(FormatStrFormatter("%.1f"))

    fig.tight_layout()

    # Animation state: current frame line/scatter + end-effector trail
    robot_line, = ax.plot([], [], [], "-", linewidth=2.0, alpha=0.95)
    robot_scatter = ax.scatter([], [], [], s=26, depthshade=True)
    trail_line, = ax.plot([], [], [], "-", linewidth=0.8, alpha=0.35, color="gray")

    ee_trail_x, ee_trail_y, ee_trail_z = [], [], []
    n_freeze = 5  # freeze last frame for 5 extra frames

    def init():
        robot_line.set_data([], [])
        robot_line.set_3d_properties([])
        trail_line.set_data([], [])
        trail_line.set_3d_properties([])
        return robot_line, robot_scatter, trail_line

    def update(k):
        # Clamp to last real frame for freeze frames
        idx = min(k, N - 1)
        color = cmap(norm(float(time_grid[idx])))
        pts_l = all_pts_line[idx]
        pts_s = all_pts_scatter[idx]

        # Current robot pose
        robot_line.set_data(pts_l[:, 0], pts_l[:, 1])
        robot_line.set_3d_properties(pts_l[:, 2])
        robot_line.set_color(color)

        robot_scatter._offsets3d = (pts_s[:, 0], pts_s[:, 1], pts_s[:, 2])
        robot_scatter.set_color([color] * len(pts_s))

        # End-effector trail (only append for real frames)
        if k < N:
            ee_trail_x.append(pts_l[-1, 0])
            ee_trail_y.append(pts_l[-1, 1])
            ee_trail_z.append(pts_l[-1, 2])
        trail_line.set_data(ee_trail_x, ee_trail_y)
        trail_line.set_3d_properties(ee_trail_z)

        return robot_line, robot_scatter, trail_line

    anim = FuncAnimation(fig, update, frames=N + n_freeze, init_func=init,
                         interval=1000 / fps, blit=False)

    writer = FFMpegWriter(fps=fps, metadata={"title": title}, bitrate=1800)
    anim.save(save_path, writer=writer, dpi=150)
    plt.close(fig)
    print(f"[animate] Saved to: {save_path}")


def save_full_report(
    out_dir,
    run_config,
    robot,
    adal,
    adal_res,
    qbase,
    adal_metrics,
    quintic_metrics,
    show_figures=False,
):
    ensure_dir(out_dir)
    fig_dir = Path(out_dir) / "figures"
    ensure_dir(fig_dir)

    fig_entries_md = []
    fig_entries_docx = []

    def add_fig(title, filename):
        relp = f"figures/{filename}"
        absp = fig_dir / filename
        fig_entries_md.append((title, relp))
        fig_entries_docx.append((title, absp))
        return str(absp)

    joint_labels = [f"joint{j+1}" for j in range(6)]

    plot_joint_bundle_single(
        adal_res["t"],
        adal_res["qdot"],
        title="ADA-L joint angular velocity",
        y_label="Angular velocity,\n$\\dot{q}$ (rad/s)",
        labels=joint_labels,
        save_path=add_fig("ADA-L joint angular velocity", "adal_joint_angular_velocity.png"),
        show=show_figures,
    )

    plot_joint_bundle_single(
        adal_res["t"],
        adal_res["qddot"],
        title="ADA-L joint angular acceleration",
        y_label="Angular acceleration,\n$\\ddot{q}$ (rad/s$^2$)",
        labels=joint_labels,
        save_path=add_fig("ADA-L joint angular acceleration", "adal_joint_angular_acceleration.png"),
        show=show_figures,
    )

    plot_joint_bundle_single(
        adal_res["t"],
        adal_res["qjerk"],
        title="ADA-L joint angular jerk",
        y_label="Angular jerk,\n$\\dddot{q}$ (rad/s$^3$)",
        labels=joint_labels,
        save_path=add_fig("ADA-L joint angular jerk", "adal_joint_angular_jerk.png"),
        show=show_figures,
    )

    plot_joint_trajectories_compare(
        adal_res["t"],
        adal_res["q"],
        qbase["q"],
        label_a="ADA-L",
        label_b="Quintic",
        save_path=add_fig("Joint trajectory comparison", "compare_joint_trajectories.png"),
        show=show_figures,
        q0=run_config.get("q0"),
    )

    plot_ee_paths_compare(
        adal_res["ee"],
        qbase["ee"],
        run_config["target_xyz"],
        label_a="ADA-L",
        label_b="Quintic",
        save_path=add_fig("End-effector path comparison", "compare_ee_paths.png"),
        show=show_figures,
        xy_center=run_config.get("xy_center"),
    )

    adal_vel_norm = np.linalg.norm(adal_res["qdot"], axis=1)
    qbase_vel_norm = np.linalg.norm(qbase["qdot"], axis=1)
    adal_acc_norm = np.linalg.norm(adal_res["qddot"], axis=1)
    qbase_acc_norm = np.linalg.norm(qbase["qddot"], axis=1)
    adal_jerk_norm = np.linalg.norm(adal_res["qjerk"], axis=1)
    qbase_jerk_norm = np.linalg.norm(qbase["qjerk"], axis=1)

    plot_norm_profiles_compare(
        adal_res["t"],
        adal_vel_norm,
        qbase_vel_norm,
        title="Angular velocity norm: ADA-L vs Quintic",
        y_label="Angular velocity \nnorm, $\\|\\dot{q}\\|$ (rad/s)",
        label_a="ADA-L",
        label_b="Quintic",
        save_path=add_fig("Angular velocity norm comparison", "compare_angular_velocity_norm.png"),
        show=show_figures,
    )

    plot_norm_profiles_compare(
        adal_res["t"],
        adal_acc_norm,
        qbase_acc_norm,
        title="Angular acceleration norm: ADA-L vs Quintic",
        y_label="Angular acceleration \nnorm, $\\|\\ddot{q}\\|$ (rad/s$^2$)",
        label_a="ADA-L",
        label_b="Quintic",
        save_path=add_fig("Angular acceleration norm comparison", "compare_angular_acceleration_norm.png"),
        show=show_figures,
    )

    plot_norm_profiles_compare(
        adal_res["t"],
        adal_jerk_norm,
        qbase_jerk_norm,
        title="Angular jerk norm: ADA-L vs Quintic",
        y_label="Angular jerk norm, \n$\\|\\dddot{q}\\|$ (rad/s$^3$)",
        label_a="ADA-L",
        label_b="Quintic",
        save_path=add_fig("Angular jerk norm comparison", "compare_angular_jerk_norm.png"),
        show=show_figures,
    )

    plot_robot_snapshots_generic(
        robot=robot,
        q=adal_res["q"],
        target_xyz=run_config["target_xyz"],
        time_grid=adal_res["t"],
        n_show=6,
        title="ADA-L trajectory snapshots",
        save_path=add_fig("ADA-L robot snapshots", "adal_robot_snapshots.png"),
        show=show_figures,
        xy_center=run_config.get("xy_center"),
    )

    plot_robot_snapshots_generic(
        robot=robot,
        q=qbase["q"],
        target_xyz=run_config["target_xyz"],
        time_grid=qbase["t"],
        n_show=6,
        title="Quintic trajectory snapshots",
        save_path=add_fig("Quintic robot snapshots", "quintic_robot_snapshots.png"),
        show=show_figures,
        xy_center=run_config.get("xy_center"),
    )

    save_json(adal_metrics, Path(out_dir) / "metrics_adal.json")
    save_json(quintic_metrics, Path(out_dir) / "metrics_quintic.json")
    save_comparison_csv(
        adal_metrics,
        quintic_metrics,
        Path(out_dir) / "metrics_comparison.csv",
        name_a="ADA-L",
        name_b="Quintic",
    )

    save_joint_trajectory_csv(adal_res["t"], adal_res["q"], Path(out_dir) / "joint_trajectory_adal.csv")
    save_joint_trajectory_csv(qbase["t"], qbase["q"], Path(out_dir) / "joint_trajectory_quintic.csv")

    np.savez(
        Path(out_dir) / "trajectory_data.npz",
        t_adal=adal_res["t"],
        q_adal=adal_res["q"],
        qdot_adal=adal_res["qdot"],
        qddot_adal=adal_res["qddot"],
        qjerk_adal=adal_res["qjerk"],
        ee_adal=adal_res["ee"],
        t_quintic=qbase["t"],
        q_quintic=qbase["q"],
        qdot_quintic=qbase["qdot"],
        qddot_quintic=qbase["qddot"],
        qjerk_quintic=qbase["qjerk"],
        ee_quintic=qbase["ee"],
        q0=run_config["q0"],
        target_xyz=run_config["target_xyz"],
    )

    write_markdown_report(
        report_path=Path(out_dir) / "report.md",
        run_config=run_config,
        metrics_adal=adal_metrics,
        metrics_quintic=quintic_metrics,
        fig_paths=fig_entries_md,
        adal_final_ee=adal_res["ee"][-1],
        quintic_final_ee=qbase["ee"][-1],
        target_xyz=run_config["target_xyz"],
    )
    print(f"[report] Markdown saved to: {Path(out_dir) / 'report.md'}")

    try_write_docx_report(
        docx_path=Path(out_dir) / "report.docx",
        run_config=run_config,
        metrics_adal=adal_metrics,
        metrics_quintic=quintic_metrics,
        fig_abs_paths=fig_entries_docx,
        adal_final_ee=adal_res["ee"][-1],
        quintic_final_ee=qbase["ee"][-1],
        target_xyz=run_config["target_xyz"],
    )

# ============================================================
# 9) Main
# ============================================================
if __name__ == "__main__":
    # --------------------------------------------------------
    # Problem setup
    # --------------------------------------------------------
    # URDF home pose -> internal q (DH theta = q + theta_offset). Match Ex1.
    q0_UR5e = np.array(
        [0, -np.pi / 2.0, np.pi / 2.0, -np.pi / 2.0, -np.pi / 2.0, np.pi / 2.0],
        dtype=np.float32,
    )
    theta_offset = np.array(
        [0, -np.pi / 2.0, 0, -np.pi / 2.0, 0, 0], dtype=np.float32,
    )
    q0 = (q0_UR5e - theta_offset).astype(np.float32)
    target_xyz = np.array([0.60, -0.15, 0.55], dtype=np.float32)

    T_final = 2.0
    Nt = 101

    robot = UR5eKinematics()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = str(Path(__file__).resolve().parent / f"results_adal_vs_quintic_{timestamp}")

    # --------------------------------------------------------
    # ADA-L solve
    # --------------------------------------------------------
    print("\nRunning ADA-L planner...")
    adal = ADALTrajectoryPlanner(
        q0=q0,
        target_xyz=target_xyz,
        T_final=T_final,
        Nt=Nt,
        gamma=1.0,
        max_order=5,
        N_p=10,
        seed=0,
    )
    adal_solve_time = adal.train(
        epochs=500,
        lr=3e-2,
        print_every=100,
        use_lbfgs=True,
        lbfgs_maxiter=2000,
    )
    adal_res = adal.results()

    adal_metrics = benchmark_metrics(
        q=adal_res["q"],
        qdot=adal_res["qdot"],
        qddot=adal_res["qddot"],
        qjerk=adal_res["qjerk"],
        ee=adal_res["ee"],
        target_xyz=adal_res["target_xyz"],
        dt=T_final / (Nt - 1),
        solve_time=adal_solve_time,
    )

    # --------------------------------------------------------
    # Quintic baseline solve
    # --------------------------------------------------------
    print("\nRunning quintic baseline...")
    quintic = JointQuinticBaseline(
        robot=robot,
        q0=q0,
        target_xyz=target_xyz,
        T_final=T_final,
        Nt=Nt,
    )
    qbase = quintic.run(
        ik_epochs=3000,
        ik_lr=3e-2,
    )

    quintic_metrics = benchmark_metrics(
        q=qbase["q"],
        qdot=qbase["qdot"],
        qddot=qbase["qddot"],
        qjerk=qbase["qjerk"],
        ee=qbase["ee"],
        target_xyz=target_xyz,
        dt=T_final / (Nt - 1),
        solve_time=qbase["solve_time"],
    )

    # --------------------------------------------------------
    # Print metrics
    # --------------------------------------------------------
    print_metrics_table("ADA-L", adal_metrics)
    print_metrics_table("Quintic baseline", quintic_metrics)
    print_comparison_table(adal_metrics, quintic_metrics, name_a="ADA-L", name_b="Quintic")

    # --------------------------------------------------------
    # Final target errors
    # --------------------------------------------------------
    print("\nFinal EE positions")
    print("ADA-L final EE      :", adal_res["ee"][-1])
    print("Quintic final EE  :", qbase["ee"][-1])
    print("Target            :", target_xyz)

    # --------------------------------------------------------
    # Save figures + report
    # With show_figures=False, only save to file and do not open a window
    # --------------------------------------------------------
    # xy_center for plot framing: midpoint of start EE and target (Ex1 pattern)
    _, _ee0 = robot.forward_all_points(tf.constant(q0[None, :], dtype=tf.float32))
    _ee0_np = _ee0.numpy()[0]
    xy_center = np.array(
        [0.5 * (float(_ee0_np[0]) + float(target_xyz[0])),
         0.5 * (float(_ee0_np[1]) + float(target_xyz[1]))],
        dtype=np.float32,
    )

    run_config = {
        "q0": q0.copy(),
        "target_xyz": target_xyz.copy(),
        "T_final": T_final,
        "Nt": Nt,
        "xy_center": xy_center,
    }

    save_full_report(
        out_dir=out_dir,
        run_config=run_config,
        robot=robot,
        adal=adal,
        adal_res=adal_res,
        qbase=qbase,
        adal_metrics=adal_metrics,
        quintic_metrics=quintic_metrics,
        show_figures=False,
    )

    print(f"\nAll outputs saved under: {out_dir}")

    # ============================================================

## ADA-L animation ##
# animate_robot_trajectory(
    # robot=robot,
    # q=adal_res["q"],
    # target_xyz=run_config["target_xyz"],
    # time_grid=adal_res["t"],
    # save_path=str(Path(out_dir) / "ADAL_animation.mp4"),
    # fps=20,
    # title="ADA-L trajectory",
    # xy_center=run_config.get("xy_center"),
# )

## Quintic animation ##
# animate_robot_trajectory(
    # robot=robot,
    # q=qbase["q"],
    # target_xyz=run_config["target_xyz"],
    # time_grid=adal_res["t"],
    # save_path=str(Path(out_dir) / "Quintic_animation.mp4"),
    # fps=20,
    # title="Quintic trajectory",
    # xy_center=run_config.get("xy_center"),
# )